"""Quality/coverage and resource-safety tests with synthetic, non-private files."""
from __future__ import annotations

import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches as WordInches
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

import extract as worker

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
CRITICAL = "Если x ≠ 0, то y ≤ −2,5; не менять условие."


class ExtractionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # A real Unicode font prevents false assertions against font substitution.
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ]
        font = next((p for p in candidates if Path(p).is_file()), None)
        if font is None:
            raise unittest.SkipTest("Unicode fixture font unavailable")
        pdfmetrics.registerFont(TTFont("FixtureUnicode", font))

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="extract-test-", dir=Path(__file__).parent)
        self.folder = Path(self.temp.name).resolve()

    def tearDown(self):
        self.temp.cleanup()

    def parse(self, path, mime):
        return worker.extract(path, self.folder / "job", mime)

    def image(self):
        path = self.folder / "fixture.png"
        image = Image.new("RGB", (220, 140), "white")
        for x in range(0, 220, 10):
            for y in range(0, 140, 10):
                image.putpixel((x, y), (x, y, 40))
        image.save(path)
        return path

    def pdf(self, *, scan=False, mixed=False, pages=1, table=False):
        path = self.folder / "fixture.pdf"
        c = canvas.Canvas(str(path))
        for _ in range(pages):
            c.setFont("FixtureUnicode", 12)
            if not scan:
                c.drawString(40, 780, CRITICAL)
                c.drawString(40, 756, "Следующий абзац: сохранить 28 октября.")
            if scan or mixed:
                c.drawImage(ImageReader(str(self.image())), 40, 380, width=330, height=210)
            if table:
                for y in (660, 630, 600):
                    c.line(40, y, 400, y)
                for x in (40, 220, 400):
                    c.line(x, 600, x, 660)
                c.drawString(50, 640, "Условие")
                c.drawString(230, 640, "Баллы")
                c.drawString(50, 612, "не менее 70%")
                c.drawString(230, 612, "−2,5")
            c.showPage()
        c.save()
        return path

    def test_pdf_preserves_critical_symbols_numbers_negation_and_order(self):
        result = self.parse(self.pdf(), "application/pdf")
        text = "\n".join(b["text"] for b in result["blocks"])
        self.assertIn(CRITICAL, text)
        self.assertLess(text.index("Если"), text.index("Следующий"))
        self.assertEqual(result["pageCount"], 1)
        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["needsVision"], [])

    def test_scanned_pdf_queues_render_with_contained_private_path(self):
        result = self.parse(self.pdf(scan=True), "application/pdf")
        self.assertEqual(result["status"], "limited")
        self.assertEqual(len(result["needsVision"]), 1)
        item = result["needsVision"][0]
        path = Path(item["imageFile"])
        self.assertTrue(path.is_relative_to(self.folder / "job"))
        self.assertEqual(path.stat().st_mode & 0o777, 0o600)
        with Image.open(path) as picture:
            self.assertLessEqual(picture.width * picture.height, 4_010_000)
        self.assertEqual(item["page"], 1)

    def test_mixed_pdf_keeps_native_text_and_visual_region_pending(self):
        result = self.parse(self.pdf(mixed=True), "application/pdf")
        self.assertIn(CRITICAL, "\n".join(b["text"] for b in result["blocks"]))
        self.assertEqual(len(result["needsVision"]), 1)
        self.assertEqual(result["status"], "limited")

    def test_pdf_ruled_table_preserves_cells_and_negation(self):
        result = self.parse(self.pdf(table=True), "application/pdf")
        tables = [b["text"] for b in result["blocks"] if b["kind"] == "table"]
        self.assertTrue(tables)
        self.assertIn("не менее 70%", tables[0])
        self.assertIn("−2,5", tables[0])

    def test_pdf_page_cap_is_explicit_not_truncated_success(self):
        path = self.pdf(pages=2)
        with patch.object(worker, "PAGE_LIMIT", 1), self.assertRaisesRegex(worker.ExtractionError, "страниц"):
            self.parse(path, "application/pdf")

    def test_recheck_renders_only_selected_page_without_reextracting_document(self):
        path = self.pdf(pages=3)
        with patch.object(worker, "extract_pdf", side_effect=AssertionError("Full extraction must not run")):
            result = worker.extract(path, self.folder / "job", "application/pdf", render_page=2)
        self.assertEqual(result["pageCount"], 3)
        self.assertEqual(len(result["blocks"]), 1)
        self.assertEqual(result["blocks"][0]["page"], 2)
        self.assertEqual(len(result["needsVision"]), 1)
        self.assertTrue(any("только выбранную" in w for w in result["warnings"]))

    def test_recheck_missing_page_or_nonpdf_is_explicit_error(self):
        path = self.pdf()
        for page in (0, 2, 301):
            with self.assertRaises(worker.ExtractionError):
                worker.extract(path, self.folder / "job", "application/pdf", render_page=page)
        txt = self.folder / "file.txt"
        txt.write_text("Text", encoding="utf-8")
        with self.assertRaisesRegex(worker.ExtractionError, "только для PDF"):
            worker.extract(txt, self.folder / "job", "text/plain", render_page=1)

    def docx(self):
        path = self.folder / "fixture.docx"
        document = Document()
        document.add_paragraph(CRITICAL)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Условие"
        table.cell(0, 1).text = "Баллы"
        table.cell(1, 0).text = "не менее 70%"
        table.cell(1, 1).text = "−2,5"
        document.add_paragraph("После таблицы")
        document.add_picture(str(self.image()), width=WordInches(1))
        document.sections[0].header.paragraphs[0].text = "Контрольный колонтитул"
        paragraph = document.add_paragraph("Формула: ")
        math = OxmlElement("m:oMath")
        fraction = OxmlElement("m:f")
        for tag, value in (("m:num", "1"), ("m:den", "x")):
            container = OxmlElement(tag)
            run = OxmlElement("m:r")
            text = OxmlElement("m:t")
            text.text = value
            run.append(text)
            container.append(run)
            fraction.append(container)
        math.append(fraction)
        paragraph._p.append(math)
        document.save(path)
        return path

    def test_docx_order_formula_header_table_and_image(self):
        result = self.parse(self.docx(), DOCX)
        text = "\n".join(b["text"] for b in result["blocks"])
        self.assertIn(CRITICAL, text)
        self.assertIn("не менее 70% | −2,5", text)
        self.assertIn("Контрольный колонтитул", text)
        self.assertIn(r"\frac{1}{x}", text)
        self.assertLess(text.index("не менее"), text.index("После таблицы"))
        self.assertEqual(len(result["needsVision"]), 1)
        self.assertEqual(result["pageCount"], 0)  # Never invent Word page numbers.
        self.assertTrue(all("page" not in block for block in result["blocks"]))

    def rewrite_zip(self, path, additions):
        target = path.with_suffix(".rewrite")
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
            for info in source.infolist():
                if info.filename not in additions:
                    output.writestr(info, source.read(info))
            for name, payload in additions.items():
                output.writestr(name, payload)
        target.replace(path)

    def test_docx_footnote_and_textbox_are_not_silently_omitted(self):
        path = self.docx()
        with zipfile.ZipFile(path) as z:
            root = worker.safe_xml(z.read("word/document.xml"))
        body = root.find("w:body", worker.NS)
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text_box = OxmlElement("w:txbxContent")
        inner_p = OxmlElement("w:p")
        inner_r = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Поле: не менять знак ≠"
        inner_r.append(text)
        inner_p.append(inner_r)
        text_box.append(inner_p)
        run.append(text_box)
        paragraph.append(run)
        body.insert(0, paragraph)
        self.rewrite_zip(path, {
            "word/document.xml": worker.etree.tostring(root),
            "word/footnotes.xml": f'<w:footnotes xmlns:w="{worker.NS["w"]}"><w:footnote w:id="2"><w:p><w:r><w:t>Сноска: только при x ≠ 0</w:t></w:r></w:p></w:footnote></w:footnotes>'.encode(),
        })
        result = self.parse(path, DOCX)
        texts = [b["text"] for b in result["blocks"]]
        self.assertEqual(sum("Поле: не менять" in t for t in texts), 1)
        self.assertTrue(any("Сноска: только при x ≠ 0" in t for t in texts))
        self.assertTrue(any("Текстовые поля" in warning for warning in result["warnings"]))

    def test_pptx_preserves_slide_order_table_notes_and_each_image_reference(self):
        path = self.folder / "fixture.pptx"
        presentation = Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = first.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text_frame
        frame.text = CRITICAL
        first.notes_slide.notes_text_frame.text = "Не забыть: 28 октября предварительно"
        table = first.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
        table.cell(0, 0).text = "Условие"
        table.cell(0, 1).text = "Баллы"
        table.cell(1, 0).text = "не менее 70%"
        table.cell(1, 1).text = "−2,5"
        first.shapes.add_picture(str(self.image()), Inches(1), Inches(4), width=Inches(1))
        first.shapes.add_picture(str(self.image()), Inches(3), Inches(4), width=Inches(1))
        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Второй слайд"
        presentation.save(path)
        result = self.parse(path, PPTX)
        text = "\n".join(b["text"] for b in result["blocks"])
        self.assertIn(CRITICAL, text)
        self.assertIn("не менее 70% | −2,5", text)
        self.assertIn("28 октября предварительно", text)
        self.assertLess(text.index("Если"), text.index("Второй слайд"))
        self.assertEqual(len(result["needsVision"]), 2)
        self.assertEqual(result["pageCount"], 2)
        self.assertEqual(len({b["blockId"] for b in result["needsVision"]}), 2)

    def test_utf8_and_utf16_preserve_source_content(self):
        for encoding in ("utf-8", "utf-16"):
            path = self.folder / f"source-{encoding}.txt"
            path.write_text(CRITICAL + "\n\nОбразец 28 октября", encoding=encoding)
            result = self.parse(path, "text/plain")
            self.assertEqual(result["status"], "ready")
            self.assertEqual(result["blocks"][0]["text"], CRITICAL)

    def test_binary_not_misreported_as_ready_text(self):
        path = self.folder / "binary.txt"
        path.write_bytes(b"\x00binary\x00")
        with self.assertRaises(worker.ExtractionError):
            self.parse(path, "text/plain")

    def test_photo_orientation_and_global_warning(self):
        path = self.folder / "phone.jpg"
        image = Image.new("RGB", (120, 60), "white")
        exif = Image.Exif()
        exif[274] = 6
        image.save(path, exif=exif)
        result = self.parse(path, "image/jpeg")
        with Image.open(result["needsVision"][0]["imageFile"]) as saved:
            self.assertEqual(saved.size, (60, 120))
        self.assertTrue(any("рекомендуется сверка" in warning for warning in result["warnings"]))

    def test_photo_mime_mismatch_rejected(self):
        with self.assertRaisesRegex(worker.ExtractionError, "формату"):
            self.parse(self.image(), "image/jpeg")

    def test_decoded_pixel_and_output_limits_fail_explicitly(self):
        path = self.image()
        with patch.object(worker, "PIXEL_LIMIT", 100), self.assertRaises(worker.ExtractionError):
            self.parse(path, "image/png")
        with patch.object(worker, "OUTPUT_LIMIT", 30), self.assertRaises(worker.ExtractionError):
            self.parse(path, "image/png")

    def test_upload_size_cap(self):
        path = self.folder / "large.txt"
        path.write_text("123456", encoding="utf-8")
        with patch.object(worker, "DOC_LIMIT", 5), self.assertRaisesRegex(worker.ExtractionError, "предел загрузки"):
            self.parse(path, "text/plain")

    def test_zip_traversal_symlink_and_duplicate_parts_rejected(self):
        for kind in ("traversal", "symlink", "duplicate"):
            path = self.folder / f"bad-{kind}.docx"
            with zipfile.ZipFile(path, "w") as z:
                if kind == "traversal":
                    z.writestr("../../outside", b"do not extract")
                elif kind == "symlink":
                    info = zipfile.ZipInfo("word/document.xml")
                    info.external_attr = (0o120777 << 16)
                    z.writestr(info, b"/etc/passwd")
                else:
                    z.writestr("word/document.xml", b"one")
                    with worker.warnings.catch_warnings():
                        worker.warnings.simplefilter("ignore", UserWarning)
                        z.writestr("word/document.xml", b"two")
            with self.assertRaisesRegex(worker.ExtractionError, "Небезопасная"):
                worker.checked_archive(path)
        self.assertFalse((self.folder.parent / "outside").exists())

    def test_zip_expansion_and_xml_entity_caps(self):
        path = self.folder / "bomb.docx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"0" * (2 * worker.MIB))
        with self.assertRaisesRegex(worker.ExtractionError, "степень сжатия"):
            worker.checked_archive(path)
        path = self.docx()
        self.rewrite_zip(path, {"unused.xml": b'<!DOCTYPE x [<!ENTITY evil SYSTEM "file:///etc/passwd">]><x>&evil;</x>'})
        with self.assertRaisesRegex(worker.ExtractionError, "XML-сущности"):
            self.parse(path, DOCX)

    def test_utf16_xml_dtd_is_rejected_before_office_parser(self):
        path = self.docx()
        xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'
        self.rewrite_zip(path, {"unused.xml": xml.encode("utf-16")})
        with self.assertRaisesRegex(worker.ExtractionError, "XML-сущности"):
            self.parse(path, DOCX)

    def test_output_symlink_rejected(self):
        path = self.folder / "source.txt"
        path.write_text("Content", encoding="utf-8")
        target = self.folder / "other"
        target.mkdir()
        link = self.folder / "linked"
        link.symlink_to(target, target_is_directory=True)
        with self.assertRaisesRegex(worker.ExtractionError, "безопасные"):
            worker.extract(path, link, "text/plain")

    def test_legacy_missing_converter_is_explicit_limited(self):
        path = self.folder / "legacy.doc"
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"dummy")
        with patch.object(worker.shutil, "which", return_value=None):
            result = self.parse(path, "application/msword")
        self.assertEqual(result["status"], "limited")
        self.assertEqual(result["blocks"], [])
        self.assertTrue(any("LibreOffice" in warning for warning in result["warnings"]))

    def test_legacy_conversion_timeout_is_an_error(self):
        path = self.folder / "legacy.doc"
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"dummy")
        with patch.object(worker.shutil, "which", return_value="soffice"), \
             patch.object(worker.subprocess, "run", side_effect=subprocess.TimeoutExpired("soffice", 45)), \
             self.assertRaisesRegex(worker.ExtractionError, "время"):
            self.parse(path, "application/msword")

    def test_cli_stdout_is_single_safe_json_and_exit_code_reports_error(self):
        path = self.folder / "sensitive-filename.txt"
        path.write_bytes(b"\x00SECRET_DOCUMENT_TEXT\x00")
        run = subprocess.run([sys.executable, str(Path(__file__).with_name("extract.py")),
                              "--input", str(path), "--output-dir", str(self.folder / "job"),
                              "--mime", "text/plain"], capture_output=True, text=True, check=False)
        self.assertEqual(run.returncode, 1)
        self.assertEqual(len(run.stdout.strip().splitlines()), 1)
        self.assertEqual(json.loads(run.stdout)["status"], "limited")
        self.assertNotIn("SECRET_DOCUMENT_TEXT", run.stdout + run.stderr)
        self.assertNotIn("sensitive-filename", run.stdout + run.stderr)


if __name__ == "__main__":
    unittest.main()
