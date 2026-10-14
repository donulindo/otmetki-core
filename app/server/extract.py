#!/usr/bin/env python3
"""Bounded, offline document extraction. stdout is one JSON object, never a log.

Images are queued for the separate vision worker, not described or guessed here.
The caller must additionally run this command with CPU/RAM and wall-time limits.
No original is overwritten; generated image names never use archive/user names.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import logging
import math
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import warnings
import zipfile
from dataclasses import dataclass

from lxml import etree
from PIL import Image, ImageOps

MIB = 1024 * 1024
DOC_LIMIT = 100 * MIB
PHOTO_LIMIT = 40 * MIB
PAGE_LIMIT = 300
ARCHIVE_LIMIT = 300 * MIB
ENTRY_LIMIT = 50 * MIB
XML_LIMIT = 20 * MIB
PIXEL_LIMIT = 24_000_000
TOTAL_PIXEL_LIMIT = 180_000_000
OUTPUT_LIMIT = 256 * MIB
IMAGE_LIMIT = 200
TEXT_LIMIT = 4_000_000
PENDING_IMAGE_WARNING = "Изображение ожидает распознавания; содержимое пока не извлечено."
PENDING_PAGE_WARNING = "Страница содержит скан или визуальные сведения; нужно распознавание изображения."
PENDING_PHOTO_WARNING = "Фотография ожидает распознавания. Для рукописи рекомендуется сверка с оригиналом."
Image.MAX_IMAGE_PIXELS = PIXEL_LIMIT

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
R_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class ExtractionError(Exception):
    """Only intentionally safe, source-independent messages reach stdout."""


@dataclass
class Budget:
    pixels: int = 0
    output_bytes: int = 0
    images: int = 0
    text_chars: int = 0


class Result:
    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.blocks: list[dict] = []
        self.warnings: list[str] = []
        self.needs_vision: list[dict] = []
        self.page_count = 0
        self.budget = Budget()

    def warn(self, message: str):
        if message not in self.warnings:
            self.warnings.append(message)

    def add(self, text: str, kind="text", *, warning: str | None = None, **location):
        if not text.strip() and kind != "image":
            return None
        self.budget.text_chars += len(text)
        if self.budget.text_chars > TEXT_LIMIT:
            raise ExtractionError("Превышен предел извлечённого текста. Разделите документ на части.")
        block = {"id": f"b{len(self.blocks)+1:05d}", "text": text, "kind": kind, **location}
        if warning:
            block["warning"] = warning
            self.warn(warning)
        self.blocks.append(block)
        return block

    def save_image(self, image: Image.Image) -> str:
        width, height = image.size
        if width <= 0 or height <= 0 or width * height > PIXEL_LIMIT:
            raise ExtractionError("Изображение слишком большое после декодирования (максимум 24 Мп).")
        self.budget.pixels += width * height
        self.budget.images += 1
        if self.budget.pixels > TOTAL_PIXEL_LIMIT or self.budget.images > IMAGE_LIMIT:
            raise ExtractionError("Превышен предел обработки изображений. Разделите документ на части.")
        data = io.BytesIO()
        image.convert("RGB").save(data, format="PNG")
        payload = data.getvalue()
        self.budget.output_bytes += len(payload)
        if self.budget.output_bytes > OUTPUT_LIMIT:
            raise ExtractionError("Превышен предел размера подготовленных изображений.")
        name = f"image-{hashlib.sha256(payload).hexdigest()[:24]}.png"
        destination = self.output_dir / name
        # No user-provided filename is used; refuse symlinks even in a reused job.
        if destination.is_symlink():
            raise ExtractionError("Небезопасный каталог результата.")
        if not destination.exists():
            fd = os.open(destination, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
            with os.fdopen(fd, "wb") as target:
                target.write(payload)
        return str(destination)

    def image(self, image: Image.Image, *, warning=None, **location):
        filename = self.save_image(image)
        message = warning or PENDING_IMAGE_WARNING
        block = self.add("", "image", warning=message, imageFile=filename, **location)
        self.needs_vision.append({"blockId": block["id"], "imageFile": filename, **location})
        return block

    def image_bytes(self, payload: bytes, *, expected_format=None, **location):
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(payload)) as source:
                    if expected_format and source.format != expected_format:
                        raise ExtractionError("Содержимое фото не соответствует выбранному формату.")
                    if getattr(source, "n_frames", 1) != 1:
                        raise ExtractionError("Многостраничное или анимированное изображение нужно разделить на отдельные фото.")
                    if source.width * source.height > PIXEL_LIMIT:
                        raise ExtractionError("Изображение слишком большое после декодирования (максимум 24 Мп).")
                    source.load()
                    oriented = ImageOps.exif_transpose(source)
                    return self.image(oriented, **location)
        except ExtractionError:
            raise
        except (Image.DecompressionBombWarning, Image.DecompressionBombError):
            raise ExtractionError("Изображение слишком большое после декодирования (максимум 24 Мп).") from None
        except Exception:
            self.warn("Встроенное изображение не удалось прочитать; проверьте оригинал.")
            self.add("[Непрочитанное изображение]", warning=self.warnings[-1], **location)
            return None

    def payload(self):
        return {
            "blocks": self.blocks,
            "warnings": self.warnings,
            "pageCount": self.page_count,
            "status": "limited" if self.warnings or self.needs_vision else "ready",
            "needsVision": self.needs_vision,
        }


def safe_xml(data: bytes):
    if len(data) > XML_LIMIT:
        raise ExtractionError("XML-часть документа превышает предел обработки.")
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ExtractionError("Документ содержит неподдерживаемые XML-сущности.")
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False, recover=False)
    try:
        root = etree.fromstring(data, parser)
        # Covers UTF-16 XML as well as the ASCII precheck, without resolving DTD.
        if root.getroottree().docinfo.doctype:
            raise ExtractionError("Документ содержит неподдерживаемые XML-сущности.")
        return root
    except etree.XMLSyntaxError:
        raise ExtractionError("Повреждена XML-структура документа.") from None


def checked_archive(path: Path) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(path)
        entries = archive.infolist()
        if len(entries) > 10_000:
            raise ExtractionError("В документе слишком много частей.")
        expanded = 0
        names = set()
        for info in entries:
            name = info.filename
            posix = PurePosixPath(name)
            mode = info.external_attr >> 16
            if ("\\" in name or posix.is_absolute() or ".." in posix.parts
                    or "\x00" in name or re.match(r"^[A-Za-z]:", name)
                    or stat.S_ISLNK(mode) or name in names):
                raise ExtractionError("Небезопасная структура архива документа.")
            names.add(name)
            if info.flag_bits & 1:
                raise ExtractionError("Документы с паролем пока не поддерживаются.")
            expanded += info.file_size
            if info.file_size > ENTRY_LIMIT or expanded > ARCHIVE_LIMIT:
                raise ExtractionError("Распакованное содержимое слишком большое. Разделите документ.")
            if info.file_size > MIB and info.file_size / max(1, info.compress_size) > 500:
                raise ExtractionError("Архив превышает безопасную степень сжатия.")
        # Finish structural/size checks before any decompression or XML parsing.
        for info in entries:
            if info.filename.endswith((".xml", ".rels")):
                if info.file_size > XML_LIMIT:
                    raise ExtractionError("XML-часть документа превышает предел обработки.")
                xml = archive.read(info)
                safe_xml(xml)
        return archive
    except ExtractionError:
        if "archive" in locals():
            archive.close()
        raise
    except (zipfile.BadZipFile, OSError):
        raise ExtractionError("Не удалось прочитать архив документа.") from None


def zip_xml(archive, name):
    try:
        info = archive.getinfo(name)
        if info.file_size > XML_LIMIT:
            raise ExtractionError("XML-часть документа превышает предел обработки.")
        return safe_xml(archive.read(info))
    except KeyError:
        raise ExtractionError("В документе отсутствует обязательная часть.") from None


def relationships(archive, part: str):
    part_path = PurePosixPath(part)
    rel_name = str(part_path.parent / "_rels" / (part_path.name + ".rels"))
    if rel_name not in archive.namelist():
        return {}
    root = zip_xml(archive, rel_name)
    result = {}
    for rel in root:
        if rel.get("TargetMode") == "External":
            continue
        target = rel.get("Target", "")
        # Internal OOXML links commonly use ../media: resolve, never extract.
        parts = list(part_path.parent.parts)
        if target.startswith("/"):
            parts = []
        for piece in target.split("/"):
            if piece in ("", "."):
                continue
            if piece == "..":
                if not parts:
                    raise ExtractionError("Небезопасная связь внутри документа.")
                parts.pop()
            else:
                parts.append(piece)
        resolved = "/".join(parts)
        if resolved in archive.namelist():
            result[rel.get("Id")] = resolved
    return result


def omml_text(node) -> tuple[str, bool]:
    """Preserve common OMML math as LaTeX-like text; never invent a layout."""
    tag = etree.QName(node).localname
    children = list(node)
    unsupported = False

    def part(name, fallback=""):
        match = node.find(f"m:{name}", NS)
        return convert(match) if match is not None else fallback

    def convert(element):
        nonlocal unsupported
        text, uncertain = omml_text(element)
        unsupported |= uncertain
        return text

    if tag in ("t",):
        return node.text or "", False
    if tag.endswith("Pr") or tag in ("ctrlPr",):
        return "", False
    if tag == "f":
        text = rf"\frac{{{part('num')}}}{{{part('den')}}}"
    elif tag == "sSup":
        text = f"{{{part('e')}}}^{{{part('sup')}}}"
    elif tag == "sSub":
        text = f"{{{part('e')}}}_{{{part('sub')}}}"
    elif tag == "sSubSup":
        text = f"{{{part('e')}}}_{{{part('sub')}}}^{{{part('sup')}}}"
    elif tag == "rad":
        degree = part("deg")
        text = rf"\sqrt[{degree}]{{{part('e')}}}" if degree else rf"\sqrt{{{part('e')}}}"
    elif tag == "d":
        props = node.find("m:dPr", NS)
        begin, end = "(", ")"
        if props is not None:
            b, e = props.find("m:begChr", NS), props.find("m:endChr", NS)
            begin = b.get(f"{{{NS['m']}}}val", "(") if b is not None else "("
            end = e.get(f"{{{NS['m']}}}val", ")") if e is not None else ")"
        text = begin + "; ".join(convert(e) for e in children if etree.QName(e).localname == "e") + end
    elif tag == "eqArr":
        text = " ; ".join(convert(e) for e in children if etree.QName(e).localname == "e")
    elif tag == "nary":
        props = node.find("m:naryPr", NS)
        char = props.find("m:chr", NS) if props is not None else None
        op = char.get(f"{{{NS['m']}}}val", "∑") if char is not None else "∑"
        text = f"{op}_{{{part('sub')}}}^{{{part('sup')}}} {part('e')}"
    else:
        text = "".join(convert(c) for c in children)
        supported = {"oMath", "oMathPara", "r", "e", "num", "den", "sub", "sup", "deg", "box"}
        if tag not in supported:
            unsupported = True
    return text, unsupported


def word_paragraph(node):
    pieces = []
    uncertain = False
    has_math = False

    def visit(el):
        nonlocal uncertain, has_math
        tag = etree.QName(el).localname
        namespace = etree.QName(el).namespace
        if el is not node and tag in ("p", "txbxContent") and namespace == NS["w"]:
            return
        if namespace == NS["m"] and tag in ("oMath", "oMathPara"):
            math_text, unsure = omml_text(el)
            pieces.append(f"${math_text}$")
            uncertain |= unsure
            has_math = True
            return
        if namespace == NS["w"]:
            if tag == "t":
                pieces.append(el.text or "")
                return
            if tag == "tab":
                pieces.append("\t")
            elif tag in ("br", "cr"):
                pieces.append("\n")
            elif tag == "noBreakHyphen":
                pieces.append("\u2011")
            elif tag == "softHyphen":
                pieces.append("\u00ad")
            elif tag == "sym":
                value = el.get(f"{{{NS['w']}}}char", "")
                pieces.append(f"[символ {value}]")
                uncertain = True
            elif tag in ("footnoteReference", "endnoteReference"):
                pieces.append(f"[сноска {el.get('{'+NS['w']+'}id', '')}]")
            elif tag in ("del", "moveFrom"):
                return  # Current document only; removed text is not the current version.
        for child in el:
            visit(child)

    visit(node)
    return "".join(pieces), has_math, uncertain


def extract_docx(path, result):
    from docx import Document
    with checked_archive(path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ExtractionError("Файл не является документом Word DOCX.")
        document = Document(path)
        # Use python-docx's parsed body and additional OOXML parts explicitly.
        body = document.element.body
        consumed_media = set()

        def process_container(container, part, section=None):
            rels = relationships(archive, part)
            for element in container:
                tag = etree.QName(element).localname
                loc = {"section": section} if section else {}
                if tag == "p":
                    text, has_math, uncertain = word_paragraph(element)
                    note = "Формула Word извлечена частично; сверьте её с оригиналом." if uncertain else None
                    result.add(text, "formula" if has_math else "text", warning=note, **loc)
                    for text_box in element.findall(".//w:txbxContent", NS):
                        result.warn("Текстовые поля Word извлечены; их положение в макете нужно сверить с оригиналом.")
                        process_container(text_box, part, section="Текстовое поле")
                    ids = [el.get(f"{{{NS['r']}}}embed") for el in element.findall(".//a:blip", NS)]
                    ids += [el.get(f"{{{NS['r']}}}id") for el in element.findall(".//v:imagedata", NS)]
                    for rid in ids:
                        media = rels.get(rid)
                        if media:
                            consumed_media.add(media)
                            result.image_bytes(archive.read(media), **loc)
                        else:
                            result.warn("Связанное изображение Word недоступно; проверьте оригинал.")
                elif tag == "tbl":
                    rows = []
                    for row in element.findall("w:tr", NS):
                        cells = []
                        for cell in row.findall("w:tc", NS):
                            paragraphs = [word_paragraph(p) for p in cell.findall("w:p", NS)]
                            texts = [item[0] for item in paragraphs]
                            if any(item[2] for item in paragraphs):
                                result.warn("Формула в таблице Word извлечена частично; сверьте оригинал.")
                            cells.append("\n".join(t for t in texts if t))
                        rows.append(" | ".join(cells))
                    result.add("\n".join(rows), "table", **loc)
                    if element.findall(".//w:tbl/w:tr/w:tc/w:tbl", NS):
                        result.warn("Вложенные таблицы Word требуют проверки структуры в оригинале.")
                    for cell in element.findall("w:tr/w:tc", NS):
                        for nested in cell.findall("w:tbl", NS):
                            process_container([nested], part, section="Вложенная таблица")
                        for p in cell.findall("w:p", NS):
                            # Text is already part of the table; process drawings only.
                            for el in p.findall(".//a:blip", NS):
                                media = rels.get(el.get(f"{{{NS['r']}}}embed"))
                                if media:
                                    consumed_media.add(media)
                                    result.image_bytes(archive.read(media), **loc)
                elif tag in ("sdt", "sdtContent", "customXml"):
                    content = element.find("w:sdtContent", NS) if tag == "sdt" else element
                    if content is not None:
                        process_container(content, part, section)
                elif tag == "altChunk":
                    result.warn("Внешняя вставка Word не извлечена; проверьте оригинал.")

        process_container(body, "word/document.xml")
        for name in archive.namelist():
            if re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                root = zip_xml(archive, name)
                process_container(root, name, "Колонтитул")
            elif name in ("word/footnotes.xml", "word/endnotes.xml"):
                root = zip_xml(archive, name)
                for note in root:
                    note_id = note.get(f"{{{NS['w']}}}id", "")
                    if note_id not in ("-1", "0"):
                        process_container(note, name, f"Сноска {note_id}")
        for name in archive.namelist():
            if name.startswith("word/media/") and not name.endswith("/") and name not in consumed_media:
                result.warn("Встроенный визуальный объект Word не связан с текстовым блоком; проверьте его положение.")
                result.image_bytes(archive.read(name), section="Визуальный объект")
        if any(name.startswith("word/charts/") and name.endswith(".xml") for name in archive.namelist()):
            result.warn("Диаграммы Word пока не извлекаются полностью; нужен просмотр оригинала.")
        if b"AlternateContent" in archive.read("word/document.xml"):
            result.warn("Альтернативные визуальные объекты Word требуют сверки с оригиналом.")


def extract_pptx(path, result):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    with checked_archive(path) as archive:
        if "ppt/presentation.xml" not in archive.namelist():
            raise ExtractionError("Файл не является презентацией PPTX.")
        presentation = Presentation(path)
        if len(presentation.slides) > PAGE_LIMIT:
            raise ExtractionError("В презентации более 300 слайдов. Разделите её на части.")
        result.page_count = len(presentation.slides)

        def shapes(items, slide):
            # XML shape order is stable; layout order cannot always be inferred.
            for shape in items:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, slide)
                    continue
                if shape.has_text_frame:
                    text = "\n".join(p.text for p in shape.text_frame.paragraphs)
                    result.add(text, slide=slide)
                if shape.has_table:
                    text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                    result.add(text, "table", slide=slide)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    result.image_bytes(shape.image.blob, slide=slide)
                if shape.has_chart:
                    result.warn(f"Диаграмма на слайде {slide} требует просмотра оригинала.")
                    result.add("[Диаграмма не извлечена полностью]", warning=result.warnings[-1], slide=slide)
                # OOXML formula objects are not included in python-pptx text.
                for math_node in shape.element.findall(".//m:oMath", NS):
                    text, uncertain = omml_text(math_node)
                    result.add(f"${text}$", "formula", slide=slide,
                               warning="Формулу на слайде нужно сверить с оригиналом." if uncertain else None)
                if shape.shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                    result.warn(f"Вложенный объект на слайде {slide} не извлечён; проверьте оригинал.")

        for index, slide in enumerate(presentation.slides, 1):
            shapes(slide.shapes, index)
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None:
                    result.add(frame.text, slide=index, section="Заметки к слайду")
        if any(name.startswith("ppt/diagrams/") and name.endswith(".xml") for name in archive.namelist()):
            result.warn("SmartArt и сложные схемы пока требуют просмотра оригинала.")


def pdf_has_images(resources, seen=None, depth=0):
    if depth > 8:
        return True
    seen = seen or set()
    if not resources:
        return False
    resources = resources.get_object()
    objects = resources.get("/XObject", {})
    objects = objects.get_object() if hasattr(objects, "get_object") else objects
    if len(objects) > IMAGE_LIMIT:
        raise ExtractionError("В PDF слишком много объектов на странице.")
    for value in objects.values():
        obj = value.get_object()
        identity = id(obj)
        if identity in seen:
            continue
        seen.add(identity)
        if obj.get("/Subtype") == "/Image":
            return True
        if obj.get("/Subtype") == "/Form" and pdf_has_images(obj.get("/Resources"), seen, depth + 1):
            return True
    return False


def render_pdf_page_handle(document, page_number, result):
    render_page = document[page_number - 1]
    try:
        width, height = render_page.get_size()
        if not math.isfinite(width * height) or width <= 0 or height <= 0:
            raise ExtractionError("Некорректный размер страницы PDF.")
        scale = min(1.8, math.sqrt(4_000_000 / (width * height)))
        if scale <= 0 or width * height > 200_000_000:
            raise ExtractionError("Размер страницы PDF превышает безопасный предел.")
        bitmap = render_page.render(scale=scale)
        try:
            result.image(bitmap.to_pil(), page=page_number, warning=PENDING_PAGE_WARNING)
        finally:
            bitmap.close()
    finally:
        render_page.close()


def render_selected_pdf_page(path, result, page_number):
    import pypdfium2 as pdfium
    if page_number < 1 or page_number > PAGE_LIMIT:
        raise ExtractionError("Номер страницы вне допустимого диапазона 1–300.")
    with pdfium.PdfDocument(str(path)) as document:
        if len(document) > PAGE_LIMIT:
            raise ExtractionError("В PDF более 300 страниц. Разделите его на части.")
        if page_number > len(document):
            raise ExtractionError("В PDF нет выбранной страницы.")
        result.page_count = len(document)
        render_pdf_page_handle(document, page_number, result)
        result.warn("Перепроверка охватывает только выбранную страницу, а не весь документ.")


def extract_pdf(path, result):
    from pypdf import PdfReader, overwrite_configuration
    import pdfplumber
    import pypdfium2 as pdfium
    # Guard common Flate decompression before text extraction, not after it.
    overwrite_configuration(
        zlib_maximum_output_length=XML_LIMIT,
        array_based_stream_maximum_output_length=XML_LIMIT,
        maximum_declared_stream_length=XML_LIMIT,
        lzw_maximum_output_length=XML_LIMIT,
        run_length_maximum_output_length=XML_LIMIT,
        brotli_maximum_output_length=XML_LIMIT,
        stream_decoding_work_maximum_length=80 * MIB,
        page_tree_maximum_entries=5000,
        xform_maximum_invocations_per_extraction=1000,
    )
    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        raise ExtractionError("PDF с паролем пока не поддерживается.")
    if len(reader.pages) > PAGE_LIMIT:
        raise ExtractionError("В PDF более 300 страниц. Разделите его на части.")
    result.page_count = len(reader.pages)
    rendered = None
    plumber = None
    try:
        for index, page in enumerate(reader.pages, 1):
            raw_text = ""
            native_error = False
            try:
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > XML_LIMIT:
                    raise ExtractionError("Поток содержимого PDF слишком большой. Разделите документ.")
                raw_text = page.extract_text() or ""
                result.add(raw_text, page=index)
            except ExtractionError:
                raise
            except Exception:
                native_error = True
                result.warn(f"Текст страницы {index} не удалось прочитать программно.")
            has_image = pdf_has_images(page.get("/Resources"))
            broken = "\ufffd" in raw_text or raw_text.count("\x00") > 0
            # Keep existing text on mixed pages and queue the full visual page.
            needs_image = has_image or not raw_text.strip() or native_error or broken
            if needs_image:
                if rendered is None:
                    rendered = pdfium.PdfDocument(str(path))
                render_pdf_page_handle(rendered, index, result)
            # Ruled tables only, without reparsing image-only pages unnecessarily.
            if raw_text.strip():
                if plumber is None:
                    plumber = pdfplumber.open(path)
                plumber_page = plumber.pages[index - 1]
                try:
                    if len(plumber_page.edges) >= 4:
                        for table in plumber_page.extract_tables():
                            if table and len(table) > 1:
                                result.add("\n".join(" | ".join(str(c or "") for c in row) for row in table),
                                           "table", page=index)
                except Exception:
                    result.warn(f"Таблицы страницы {index} требуют сверки с оригиналом.")
                finally:
                    plumber_page.close()
    finally:
        if plumber is not None:
            plumber.close()
        if rendered is not None:
            rendered.close()


def extract_txt(path, result):
    from charset_normalizer import from_bytes
    raw = path.read_bytes()
    if len(raw) > TEXT_LIMIT * 4:
        raise ExtractionError("Текстовый файл слишком большой для одной обработки.")
    try:
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = raw.decode("utf-16", errors="strict")
        else:
            text = raw.decode("utf-8-sig", errors="strict")
    except UnicodeError:
        match = from_bytes(raw).best()
        if match is None or match.chaos > 0.1:
            raise ExtractionError("Кодировка текста неоднозначна. Сохраните файл в UTF-8.")
        text = str(match)
        result.warn(f"Кодировка определена автоматически ({match.encoding}); рекомендуется сверка текста.")
    if "\x00" in text or "\ufffd" in text:
        raise ExtractionError("Файл не является корректным текстом. Сохраните его в UTF-8.")
    # Preserve paragraphs; no whitespace/sign normalization during extraction.
    for paragraph in re.split(r"\r?\n\s*\r?\n", text):
        result.add(paragraph)


def convert_legacy(path, result, format_name):
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        result.warn("Для старого формата нужна конвертация LibreOffice. Загрузите DOCX или PPTX.")
        return
    with tempfile.TemporaryDirectory(prefix="office-", dir=result.output_dir) as temp:
        work = Path(temp)
        # Copies get a fixed name and extension: no flag-like user names.
        source = work / ("source.doc" if format_name == "docx" else "source.ppt")
        shutil.copyfile(path, source)
        profile = work / "profile"
        try:
            run = subprocess.run(
                [binary, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--nologo",
                 "--nodefault", "--nofirststartwizard", "--convert-to", format_name,
                 "--outdir", str(work), str(source)],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=45, check=False,
            )
        except subprocess.TimeoutExpired:
            raise ExtractionError("Конвертация Office превысила допустимое время.") from None
        target = work / f"source.{format_name}"
        if run.returncode != 0 or not target.is_file() or target.is_symlink():
            raise ExtractionError("Не удалось преобразовать старый формат Office.")
        if target.stat().st_size > DOC_LIMIT:
            raise ExtractionError("Преобразованный документ превышает предел размера.")
        if format_name == "docx":
            extract_docx(target, result)
        else:
            extract_pptx(target, result)
        result.warn("Старый формат преобразован; оформление и полноту следует сверить с оригиналом.")


def extract(input_path: str | Path, output_dir: str | Path, mime: str, render_page: int | None = None) -> dict:
    path = Path(input_path)
    output = Path(output_dir)
    if not path.is_absolute() or not output.is_absolute() or path.is_symlink() or output.is_symlink():
        raise ExtractionError("Для обработки нужны безопасные абсолютные пути.")
    if not path.is_file():
        raise ExtractionError("Исходный файл недоступен.")
    size = path.stat().st_size
    is_photo = mime in ("image/png", "image/jpeg", "image/webp")
    if size > (PHOTO_LIMIT if is_photo else DOC_LIMIT):
        raise ExtractionError("Файл превышает предел загрузки (100 МиБ документ, 40 МиБ фото).")
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    output = output.resolve()
    result = Result(output)
    with path.open("rb") as source:
        prefix = source.read(16)
    if mime == "application/pdf":
        if not prefix.startswith(b"%PDF-"):
            raise ExtractionError("Содержимое файла не соответствует PDF.")
        if render_page is not None:
            render_selected_pdf_page(path, result, render_page)
        else:
            extract_pdf(path, result)
    elif render_page is not None:
        raise ExtractionError("Перепроверка страницы поддерживается только для PDF.")
    elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        if not prefix.startswith(b"PK"):
            raise ExtractionError("Содержимое файла не соответствует DOCX.")
        extract_docx(path, result)
    elif mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        if not prefix.startswith(b"PK"):
            raise ExtractionError("Содержимое файла не соответствует PPTX.")
        extract_pptx(path, result)
    elif mime == "text/plain":
        extract_txt(path, result)
    elif is_photo:
        payload = path.read_bytes()
        result.page_count = 1
        result.image_bytes(payload, expected_format={"image/png": "PNG", "image/jpeg": "JPEG", "image/webp": "WEBP"}[mime],
                           page=1, warning=PENDING_PHOTO_WARNING)
    elif mime in ("application/msword", "application/vnd.ms-powerpoint"):
        if not prefix.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ExtractionError("Содержимое файла не соответствует старому формату Office.")
        convert_legacy(path, result, "docx" if mime == "application/msword" else "pptx")
    else:
        raise ExtractionError("Этот формат пока не поддерживается. Используйте PDF, DOCX, PPTX, TXT или фото.")
    if not result.blocks and not result.warnings:
        result.warn("Документ не содержит доступного текста; проверьте оригинал.")
    return result.payload()


def main(argv=None):
    parser = argparse.ArgumentParser(description="Offline bounded document extraction")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--mime", required=True)
    parser.add_argument("--render-page", type=int, help="Render one PDF page for a targeted recheck, 1-based")
    args = parser.parse_args(argv)
    # Suppress parser/library diagnostics: source text and object names are private.
    logging.disable(logging.CRITICAL)
    try:
        if sys.platform.startswith("linux"):
            import resource
            # The parent still supplies a wall-clock timeout and container CPU caps.
            for limit, requested in ((resource.RLIMIT_AS, 3 * 1024 ** 3), (resource.RLIMIT_CPU, 120)):
                soft, hard = resource.getrlimit(limit)
                bounded = requested if hard == resource.RLIM_INFINITY else min(requested, hard)
                resource.setrlimit(limit, (bounded, bounded))
        payload = extract(args.input, args.output_dir, args.mime, render_page=args.render_page)
        exit_code = 0
    except ExtractionError as error:
        payload = {"blocks": [], "warnings": [], "pageCount": 0, "status": "limited", "needsVision": [], "error": str(error)}
        exit_code = 1
    except Exception:
        payload = {"blocks": [], "warnings": [], "pageCount": 0, "status": "limited", "needsVision": [], "error": "Не удалось прочитать документ. Проверьте формат или загрузите другую версию."}
        exit_code = 1
    print(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
